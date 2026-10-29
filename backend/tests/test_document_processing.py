import io
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from pathlib import Path
import tempfile
import shutil
from pypdf import PdfWriter
from docx import Document
from PIL import Image

from app.main import app
from app.api.deps import get_repository
from app.repositories.sqlite import SQLiteWorkspaceRepository
from app.config import settings


@pytest.fixture(scope="session")
def temp_test_env():
    temp_dir = tempfile.mkdtemp()
    data_dir = Path(temp_dir) / "data"
    data_dir.mkdir(parents=True, exist_ok=True)

    # Patch settings to use temp data directory
    original_data_dir = settings.DATA_DIR
    settings.DATA_DIR = data_dir

    yield data_dir

    settings.DATA_DIR = original_data_dir
    shutil.rmtree(temp_dir, ignore_errors=True)


@pytest_asyncio.fixture
async def client(temp_test_env):
    test_db_path = temp_test_env / "test_afg_phase3.db"
    test_repo = SQLiteWorkspaceRepository(test_db_path)
    await test_repo.init_db()

    app.dependency_overrides[get_repository] = lambda: test_repo

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as ac:
        yield ac

    app.dependency_overrides.clear()


def create_dummy_pdf_bytes(text: str = "Binary Search Algorithm divides the array in half.") -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    # We can write a stream or text
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_dummy_docx_bytes(text: str = "Stack and Queue Data Structure Notes") -> bytes:
    doc = Document()
    doc.add_heading("Unit 2: Linear Data Structures", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_dummy_png_bytes() -> bytes:
    img = Image.new("RGB", (60, 30), color=(73, 109, 137))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.mark.asyncio
async def test_txt_upload_and_page_extraction(client: AsyncClient):
    # 1. Create a subject
    subj_res = await client.post(
        "/api/subjects",
        json={"name": "Operating Systems", "code": "CS301"},
    )
    assert subj_res.status_code == 201
    subject_id = subj_res.json()["id"]

    # 2. Upload a TXT document
    txt_content = b"Process Scheduling algorithms include FCFS, SJF, and Round Robin.\nContext switching saves PCB state."
    files = {
        "file": ("scheduling_notes.txt", txt_content, "text/plain"),
    }
    data = {
        "title": "CPU Scheduling Notes",
        "source_type": "notes",
    }
    upload_res = await client.post(
        f"/api/subjects/{subject_id}/upload",
        files=files,
        data=data,
    )
    assert upload_res.status_code == 201
    res_data = upload_res.json()
    assert res_data["status"] == "ready"
    assert res_data["page_count"] == 1
    resource_id = res_data["id"]

    # 3. Retrieve extracted pages
    pages_res = await client.get(f"/api/resources/{resource_id}/pages")
    assert pages_res.status_code == 200
    pages_data = pages_res.json()
    assert pages_data["page_count"] == 1
    assert "FCFS, SJF" in pages_data["pages"][0]["content"]

    # 4. Download original file
    download_res = await client.get(f"/api/resources/{resource_id}/download")
    assert download_res.status_code == 200
    assert download_res.content == txt_content


@pytest.mark.asyncio
async def test_docx_upload_and_extraction(client: AsyncClient):
    subj_res = await client.post(
        "/api/subjects",
        json={"name": "Computer Networks", "code": "CS401"},
    )
    subject_id = subj_res.json()["id"]

    docx_bytes = create_dummy_docx_bytes("OSI Model 7 layers: Physical, Data Link, Network, Transport, Session, Presentation, Application.")
    files = {
        "file": ("osi_model.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    }
    data = {
        "title": "OSI 7 Layers Lecture Notes",
        "source_type": "notes",
    }
    upload_res = await client.post(
        f"/api/subjects/{subject_id}/upload",
        files=files,
        data=data,
    )
    assert upload_res.status_code == 201
    res_data = upload_res.json()
    assert res_data["status"] == "ready"
    assert res_data["page_count"] >= 1

    # Check pages
    pages_res = await client.get(f"/api/resources/{res_data['id']}/pages")
    assert pages_res.status_code == 200
    assert "OSI Model" in pages_res.json()["pages"][0]["content"]


@pytest.mark.asyncio
async def test_image_upload_and_metadata(client: AsyncClient):
    subj_res = await client.post(
        "/api/subjects",
        json={"name": "Digital Electronics", "code": "EC201"},
    )
    subject_id = subj_res.json()["id"]

    png_bytes = create_dummy_png_bytes()
    files = {
        "file": ("circuit_diagram.png", png_bytes, "image/png"),
    }
    data = {
        "title": "JK Flip-Flop Circuit Diagram",
        "source_type": "notes",
    }
    upload_res = await client.post(
        f"/api/subjects/{subject_id}/upload",
        files=files,
        data=data,
    )
    assert upload_res.status_code == 201
    res_data = upload_res.json()
    assert res_data["status"] == "ready"
    assert res_data["page_count"] == 1


@pytest.mark.asyncio
async def test_empty_file_failure_handling(client: AsyncClient):
    subj_res = await client.post(
        "/api/subjects",
        json={"name": "Database Systems", "code": "CS302"},
    )
    subject_id = subj_res.json()["id"]

    empty_bytes = b""
    files = {
        "file": ("empty_notes.txt", empty_bytes, "text/plain"),
    }
    data = {
        "title": "Empty Document Test",
        "source_type": "notes",
    }
    upload_res = await client.post(
        f"/api/subjects/{subject_id}/upload",
        files=files,
        data=data,
    )
    assert upload_res.status_code == 201
    res_data = upload_res.json()
    # Explicit failure visibility: status must be failed, with error message
    assert res_data["status"] == "failed"
    assert "empty" in res_data["error_message"].lower()

